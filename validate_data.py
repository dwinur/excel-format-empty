#!/usr/bin/env python3
"""
Data Validation Script
Validates extracted data and filled Excel files.
"""

import logging
from pathlib import Path
from typing import Dict, List, Any, Tuple
import openpyxl

from utils import logger


class DataValidator:
    """Validates extracted data and Excel outputs."""
    
    def __init__(self):
        """Initialize the validator."""
        self.validation_errors = []
        self.validation_warnings = []
    
    def validate_general_info(self, data: Dict[str, Any]) -> bool:
        """
        Validate general information data.
        
        Args:
            data: Dictionary with general info
            
        Returns:
            True if valid, False otherwise
        """
        required_fields = ['sample_name', 'sample_type']
        missing_fields = []
        
        for field in required_fields:
            if field not in data or data[field] is None or data[field] == '':
                missing_fields.append(field)
        
        if missing_fields:
            error = f"Missing required general info fields: {', '.join(missing_fields)}"
            self.validation_errors.append(error)
            logger.error(error)
            return False
        
        logger.info("General info validation passed")
        return True
    
    def validate_compositional_data(self, data: List[Dict[str, Any]]) -> bool:
        """
        Validate compositional data.
        
        Args:
            data: List of component dictionaries
            
        Returns:
            True if valid, False otherwise
        """
        if not data:
            error = "No compositional data found"
            self.validation_errors.append(error)
            logger.error(error)
            return False
        
        # Check for required fields in each component
        for i, component in enumerate(data):
            if 'name' not in component or not component['name']:
                error = f"Component {i} missing name"
                self.validation_errors.append(error)
                logger.error(error)
                return False
            
            # Check if at least one fraction is present
            if 'mole_fraction' not in component and 'mass_fraction' not in component:
                warning = f"Component {component['name']} has no fraction data"
                self.validation_warnings.append(warning)
                logger.warning(warning)
        
        # Validate mole fractions sum to approximately 1.0
        mole_fractions = [c.get('mole_fraction', 0) for c in data if c.get('mole_fraction')]
        if mole_fractions:
            total = sum(mole_fractions)
            if not (0.95 <= total <= 1.05):
                warning = f"Mole fractions sum to {total:.4f}, expected ~1.0"
                self.validation_warnings.append(warning)
                logger.warning(warning)
        
        logger.info(f"Compositional data validation passed ({len(data)} components)")
        return True
    
    def validate_pvt_experiment(self, data: Dict[str, Any]) -> bool:
        """
        Validate PVT experiment data.
        
        Args:
            data: Dictionary with PVT experiment info
            
        Returns:
            True if valid, False otherwise
        """
        if not data:
            warning = "No PVT experiment data found"
            self.validation_warnings.append(warning)
            logger.warning(warning)
            return True  # Not critical
        
        # Check for saturation pressure
        if 'saturation_pressure' in data:
            pressure = data['saturation_pressure']
            if pressure and (pressure < 0 or pressure > 20000):
                warning = f"Saturation pressure {pressure} is outside typical range"
                self.validation_warnings.append(warning)
                logger.warning(warning)
        
        logger.info("PVT experiment validation passed")
        return True
    
    def validate_cvd_data(self, data: List[Dict[str, Any]]) -> bool:
        """
        Validate CVD experimental data.
        
        Args:
            data: List of CVD stage dictionaries
            
        Returns:
            True if valid, False otherwise
        """
        if not data:
            warning = "No CVD data found"
            self.validation_warnings.append(warning)
            logger.warning(warning)
            return True  # Not critical
        
        # Check for monotonically decreasing pressure
        pressures = [stage.get('pressure') for stage in data if stage.get('pressure')]
        if len(pressures) > 1:
            for i in range(1, len(pressures)):
                if pressures[i] > pressures[i-1]:
                    warning = f"CVD pressure not monotonically decreasing at stage {i}"
                    self.validation_warnings.append(warning)
                    logger.warning(warning)
                    break
        
        logger.info(f"CVD data validation passed ({len(data)} stages)")
        return True
    
    def validate_extracted_data(self, data: Dict[str, Any]) -> bool:
        """
        Validate all extracted data.
        
        Args:
            data: Dictionary with all extracted data
            
        Returns:
            True if valid, False otherwise
        """
        logger.info("Starting data validation")
        self.validation_errors = []
        self.validation_warnings = []
        
        is_valid = True
        
        # Validate each section
        if 'general_info' in data:
            is_valid &= self.validate_general_info(data['general_info'])
        
        if 'compositional_data' in data:
            is_valid &= self.validate_compositional_data(data['compositional_data'])
        
        if 'pvt_experiment' in data:
            is_valid &= self.validate_pvt_experiment(data['pvt_experiment'])
        
        if 'cvd_data' in data:
            is_valid &= self.validate_cvd_data(data['cvd_data'])
        
        # Log summary
        logger.info(f"Validation complete: {len(self.validation_errors)} errors, "
                   f"{len(self.validation_warnings)} warnings")
        
        return is_valid
    
    def validate_excel_output(self, excel_path: Path, 
                             expected_sheets: List[str] = None) -> bool:
        """
        Validate the output Excel file.
        
        Args:
            excel_path: Path to Excel file
            expected_sheets: List of expected sheet names
            
        Returns:
            True if valid, False otherwise
        """
        if not excel_path.exists():
            error = f"Excel file not found: {excel_path}"
            self.validation_errors.append(error)
            logger.error(error)
            return False
        
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            
            # Check sheets exist
            if expected_sheets:
                missing_sheets = [s for s in expected_sheets if s not in wb.sheetnames]
                if missing_sheets:
                    error = f"Missing sheets: {', '.join(missing_sheets)}"
                    self.validation_errors.append(error)
                    logger.error(error)
                    return False
            
            # Check if file is readable
            logger.info(f"Excel file validated: {len(wb.sheetnames)} sheets found")
            wb.close()
            return True
            
        except Exception as e:
            error = f"Error validating Excel file: {str(e)}"
            self.validation_errors.append(error)
            logger.error(error)
            return False
    
    def generate_validation_report(self, output_path: Path = None) -> str:
        """
        Generate a validation report.
        
        Args:
            output_path: Optional path to save report
            
        Returns:
            Report as string
        """
        report_lines = [
            "=" * 80,
            "DATA VALIDATION REPORT",
            "=" * 80,
            ""
        ]
        
        # Errors section
        if self.validation_errors:
            report_lines.append("ERRORS:")
            for error in self.validation_errors:
                report_lines.append(f"  ❌ {error}")
            report_lines.append("")
        else:
            report_lines.append("ERRORS: None")
            report_lines.append("")
        
        # Warnings section
        if self.validation_warnings:
            report_lines.append("WARNINGS:")
            for warning in self.validation_warnings:
                report_lines.append(f"  ⚠️  {warning}")
            report_lines.append("")
        else:
            report_lines.append("WARNINGS: None")
            report_lines.append("")
        
        # Summary
        if not self.validation_errors:
            report_lines.append("✅ VALIDATION PASSED")
        else:
            report_lines.append("❌ VALIDATION FAILED")
        
        report_lines.append("=" * 80)
        
        report = "\n".join(report_lines)
        
        # Save to file if path provided
        if output_path:
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with open(output_path, 'w') as f:
                f.write(report)
            logger.info(f"Validation report saved to: {output_path}")
        
        return report


def validate_data(extracted_data: Dict[str, Any], 
                 excel_output_path: Path = None) -> Tuple[bool, str]:
    """
    Main validation function.
    
    Args:
        extracted_data: Dictionary with extracted data
        excel_output_path: Optional path to Excel output file
        
    Returns:
        Tuple of (is_valid, report)
    """
    validator = DataValidator()
    
    # Validate extracted data
    data_valid = validator.validate_extracted_data(extracted_data)
    
    # Validate Excel output if provided
    excel_valid = True
    if excel_output_path:
        excel_valid = validator.validate_excel_output(excel_output_path)
    
    # Generate report
    report = validator.generate_validation_report()
    
    is_valid = data_valid and excel_valid
    
    return is_valid, report


if __name__ == "__main__":
    # Test validation
    import sys
    
    if len(sys.argv) > 1:
        excel_path = Path(sys.argv[1])
        
        validator = DataValidator()
        is_valid = validator.validate_excel_output(
            excel_path,
            expected_sheets=['General Info', 'Compositional Data']
        )
        
        report = validator.generate_validation_report()
        print(report)
        
        sys.exit(0 if is_valid else 1)
    else:
        print("Usage: python validate_data.py <excel_path>")
