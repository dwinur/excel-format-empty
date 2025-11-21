#!/usr/bin/env python3
"""
Script to analyze the structure of Excel templates.
This helps understand the format before implementing the data mapping.
"""

import openpyxl
import sys
from pathlib import Path


def analyze_excel_structure(excel_path):
    """Analyze and print the structure of an Excel file."""
    print(f"\n{'='*80}")
    print(f"Analyzing: {excel_path}")
    print(f"{'='*80}\n")
    
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        
        print(f"Number of sheets: {len(wb.sheetnames)}")
        print(f"Sheet names: {', '.join(wb.sheetnames)}\n")
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"\n--- Sheet: {sheet_name} ---")
            print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")
            
            # Show first 10 rows with content
            print("\nFirst 10 rows with content:")
            row_count = 0
            for row in ws.iter_rows(min_row=1, max_row=20, values_only=False):
                if row_count >= 10:
                    break
                # Check if row has any content
                if any(cell.value is not None for cell in row):
                    values = []
                    for cell in row[:10]:  # First 10 columns
                        if cell.value is not None:
                            values.append(f"{cell.coordinate}: {cell.value}")
                    if values:
                        print(f"  Row {row[0].row}: {' | '.join(values)}")
                        row_count += 1
            
        wb.close()
        return True
        
    except Exception as e:
        print(f"Error analyzing {excel_path}: {str(e)}")
        return False


def main():
    """Main function to analyze Excel files."""
    current_dir = Path(__file__).parent
    
    # Target file for this phase
    target_file = "WF-Gas-Condensate-REC-Sample (CL-63169).xlsx"
    target_path = current_dir / target_file
    
    if target_path.exists():
        analyze_excel_structure(target_path)
    else:
        print(f"Target file not found: {target_path}")
        
    # Optionally analyze all Excel files
    if len(sys.argv) > 1 and sys.argv[1] == "--all":
        print("\n\nAnalyzing all Excel files:")
        for excel_file in current_dir.glob("*.xlsx"):
            if not excel_file.name.startswith("~"):  # Skip temp files
                analyze_excel_structure(excel_file)


if __name__ == "__main__":
    main()
